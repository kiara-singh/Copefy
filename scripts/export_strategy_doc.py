"""Export the Agentic RAG Testing Strategy Map to a Word document for Google Docs."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "Agentic_RAG_Testing_Strategy_Map.docx"


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            table.rows[row_idx].cells[col_idx].text = value
    doc.add_paragraph()


def main() -> None:
    doc = Document()

    title = doc.add_heading("Agentic RAG System: Testing & Strategy Map", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(
        "A planning overview for managers and partners — what we're building, "
        "what can be tested, and what's left."
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "Copefy's chat agent makes several layered decisions on every user message: "
        "which tool or capability to use, whether to search stored data, what to search "
        "(clinical content vs the user's own journal entries), how to retrieve and rank "
        "results, and how to generate a response. Each layer has different strategy options "
        "and different testing needs."
    )
    doc.add_paragraph(
        "Today, the retrieval layers (should we search? how do we search?) are the most "
        "mature and have automated evaluation — but only for the clinical knowledge base. "
        "The orchestration layer (which tool to use) and response quality layer "
        "(is the final answer actually good?) are the biggest gaps."
    )

    doc.add_heading("Data sources (important context)", level=1)
    doc.add_paragraph(
        "All searchable content lives in our database. There are two distinct corpora:"
    )
    add_table(
        doc,
        ["Corpus", "What it contains", "Example use case"],
        [
            [
                "Clinical knowledge base",
                "Curated medical and coping documents (articles, guides, strategies)",
                "User shares ADHD symptoms → retrieve relevant coping strategies",
            ],
            [
                "User journal entries",
                "The user's own past journal posts, stored per user",
                'User says "I\'m feeling anxious again" → retrieve when they last felt anxious; '
                "user mentions an exam → retrieve how they've written about exams before, their tone, patterns",
            ],
        ],
    )
    doc.add_paragraph(
        "How they differ: The clinical KB provides general knowledge (what might help). "
        "Journal entries provide personal context about this specific user — their history, "
        "emotional patterns, recurring situations, and how they tend to talk about things. "
        "Users won't usually ask \"how did I feel last week?\" directly; instead, the agent "
        "proactively pulls relevant journal context to make responses more personalized."
    )
    doc.add_paragraph(
        "Both are retrieved via semantic search over embedded content in the database. "
        "Journal search is not a separate external tool — it is part of database/KB retrieval, "
        "targeting a different corpus. The orchestrator decides whether to search, and which "
        "corpus(es) to search."
    )

    doc.add_heading("How the layers relate", level=1)
    doc.add_paragraph(
        "The orchestrator makes a single top-level decision on each message. "
        "If it chooses search database, corpus selection and retrieval strategy follow — "
        "there is no second search-or-skip step."
    )
    flowchart_png = Path(__file__).resolve().parents[1] / "docs" / "flowcharts" / "agent_flow.png"
    if flowchart_png.exists():
        doc.add_picture(str(flowchart_png), width=Inches(6.5))

    add_table(
        doc,
        ["Layer", "Question it answers", "Status"],
        [
            [
                "Agent orchestration",
                "Which path? (crisis, direct answer, search database, web, structured query, doc extraction)",
                "Partial — crisis, direct answer, and clinical KB search work; "
                "journal retrieval, web search, structured query, and document extraction are placeholders",
            ],
            [
                "Corpus selection",
                "If search database → clinical KB, journal entries, or both?",
                "Partially built — journal routing exists but retrieval not implemented",
            ],
            [
                "Retrieval strategy",
                "How do we find and rank the best content?",
                "Built for clinical KB — 4 strategies ready to compare",
            ],
            [
                "Generation",
                "How do we craft the prompt and produce the reply?",
                "Built, not yet evaluated",
            ],
            [
                "End-to-end quality",
                "Is the final answer empathetic, safe, relevant, and grounded?",
                "Not built",
            ],
        ],
    )
    doc.add_paragraph(
        "Note: Today the direct answer vs search database logic lives in separate code paths "
        "that overlap. Consolidating into one orchestrator decision is planned work."
    )

    doc.add_heading("Layer 1: Agent Orchestration", level=1)
    doc.add_paragraph(
        "What it controls: The macro routing decision — which tool or capability the agent "
        "uses before generating a reply."
    )
    doc.add_heading("Available tools", level=2)
    add_table(
        doc,
        ["Tool", "What it does", "Status"],
        [
            ["Crisis response", "Fixed crisis message (988) — no AI generation", "Built"],
            ["Direct answer", "LLM only, no tools or retrieval", "Built"],
            [
                "Database / KB search",
                "Semantic search — clinical documents and/or journal entries",
                "Partial — clinical KB works; journal retrieval is a placeholder",
            ],
            ["Web search", "External factual information from the internet", "Placeholder only"],
            ["Structured data query", "User profile, mood logs, tabular data", "Placeholder only"],
            ["Document extraction", "Parse uploaded documents", "Placeholder only"],
        ],
    )
    doc.add_heading("Orchestration strategy options", level=2)
    add_table(
        doc,
        ["Strategy", "Description", "Status"],
        [
            ["Rule-based intent router", "Pattern matching to classify intent and select tools", "Built"],
            ["AI-driven tool router", "LLM chooses which tool(s) to call", "Not built"],
            ["Multi-step agent loop", "Chain multiple tool calls", "Not built"],
        ],
    )
    doc.add_heading("What needs to be built", level=2)
    for item in [
        "Implement journal entry retrieval within database search",
        "Activate remaining placeholder tools (web search, structured query, document extraction)",
        'Unify routing logic so one system decides "which tool(s)?"',
        "Optionally: compare rule-based routing vs AI-driven tool selection",
        "Multi-step loops for queries that need chained capabilities",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("How to test", level=2)
    add_table(
        doc,
        ["Test", "What it measures", "Approach"],
        [
            ["Intent classification accuracy", "Correct intent per test case", "Labeled test suite"],
            ["Tool selection accuracy", "Correct tool(s) invoked", "Compare expected vs actual"],
            ["Crisis detection recall", "Crisis → fixed response, no AI", "Target 100%"],
            ["False tool activation", "No unnecessary tools on casual messages", "Negative test cases"],
            ["End-to-end smoke tests", "Sensible responses", "Manual and automated chat tests"],
            ["Latency by path", "Response time per routing path", "Instrument and aggregate"],
        ],
    )

    doc.add_heading("Layer 2: Direct answer vs. database search", level=1)
    doc.add_paragraph(
        "What it controls: The orchestrator's decision between LLM-only (direct answer) and "
        "searching stored data (search database). Not a separate gate after choosing database search."
    )
    doc.add_heading("Routing strategy options", level=2)
    add_table(
        doc,
        ["Strategy", "Description", "Status"],
        [
            ["Heuristic routing (current)", "Direct answer for venting/vague; search database on clinical content", "Built"],
            ["Always search database", "Search on every message", "Baseline comparison"],
            ["Never search database", "Pure conversational companion", "Baseline comparison"],
            ["AI classifier", "Model decides direct answer vs search database", "Not built"],
        ],
    )
    doc.add_heading("Corpus selection (after search database)", level=2)
    add_table(
        doc,
        ["Corpus target", "When to use", "Status"],
        [
            ["Clinical KB only", "Symptoms, coping strategies, clinical grounding", "Built"],
            ["Journal entries only", "Personal history adds context (recurring anxiety, exams) — even without referencing journal", "Routing exists; retrieval not implemented"],
            ["Both", "Clinical knowledge AND personalized context from past writing", "Not built"],
        ],
    )
    doc.add_heading("How to test", level=2)
    add_table(
        doc,
        ["Test", "Metric", "Target"],
        [
            ["Routing accuracy", "Direct answer vs search database matches expectation", "100% seeds; ≥95% full suite"],
            ["Corpus selection accuracy", "Correct corpus targeted", "Not yet testable"],
            ["Per-scenario breakdown", "Accuracy by message type", "Venting/vague must always skip"],
            ["False positive rate", "Retrieved when shouldn't", "0%"],
            ["False negative rate", "Skipped when should retrieve", "<5%"],
        ],
    )

    doc.add_heading("Layer 3: Retrieval Strategies (how to retrieve)", level=1)
    add_table(
        doc,
        ["Strategy", "Approach", "Key difference"],
        [
            ["Baseline", "Standard width + heuristic reranking", "Default"],
            ["Similarity threshold", "Drop below 0.65 confidence", "Fewer, higher-confidence chunks"],
            ["MMR (diversity)", "Balance relevance vs redundancy", "Avoids duplicate chunks"],
            ["Wide-then-narrow", "Larger pool, then narrow", "More candidates considered"],
        ],
    )
    doc.add_heading("How to test", level=2)
    add_table(
        doc,
        ["Test", "What it measures"],
        [
            ["Category recall", "Retrieved chunks match expected clinical categories"],
            ["Journal recall", "Journal entries match query (not yet testable)"],
            ["Empty retrieval rate", "How often zero results returned"],
            ["Strategy bake-off", "Head-to-head on recall, confidence, speed"],
            ["Per-scenario breakdown", "Best strategy per message type"],
        ],
    )

    doc.add_heading("Layer 4: Generation & Prompting", level=1)
    add_table(
        doc,
        ["Area", "Current approach", "Options to test"],
        [
            ["System instructions", "Single companion prompt", "Per-intent variants; stricter grounding"],
            ["Context injection", "Top 3 chunks as background", "Vary count; citations vs silent grounding"],
            ["Conversation history", "Last 6 turns", "Shorter vs longer windows"],
            ["Anti-repetition", "Hints from last reply", "On/off"],
            ["Model selection", "Fast model with fallback", "Fast vs capable model"],
            ["Generation settings", "Default config", "Lower temperature for consistency"],
        ],
    )

    doc.add_heading("Layer 5: End-to-End Response Evaluation", level=1)
    add_table(
        doc,
        ["Strategy", "What it measures", "Status"],
        [
            ["Retrieval metrics", "Gate accuracy, category recall", "Built (clinical KB)"],
            ["Rule-based checks", "Crisis keywords, banned phrases", "Not built"],
            ["AI-as-judge", "Rubric scoring", "Not built"],
            ["Pairwise comparison", "Strategy A vs B", "Not built"],
            ["Human evaluation", "Annotator scores", "Not built"],
            ["Reference answer matching", "Compare to ideal answer", "Not built"],
        ],
    )
    doc.add_heading("Proposed scoring dimensions (AI-as-judge)", level=2)
    for dim in [
        "Empathy — warm, non-judgmental tone",
        "Relevance — addresses what the user said",
        "Safety — no diagnosis, no medical advice, crisis handled correctly",
        "Grounding — aligns with sources; no fabricated facts",
        "Brevity — respects concise response guideline",
        "Non-repetition — doesn't parrot user or repeat prior message",
    ]:
        doc.add_paragraph(dim, style="List Number")

    doc.add_heading("Layer 6: Tools summary", level=1)
    add_table(
        doc,
        ["Tool", "Type", "What it does", "Status"],
        [
            ["Database search — clinical KB", "Internal retrieval", "Medical/coping documents", "Built"],
            ["Database search — journal entries", "Internal retrieval", "Personal context (patterns, history, tone)", "Placeholder"],
            ["Web search", "External tool", "Internet lookup", "Placeholder"],
            ["Structured data query", "Internal tool", "Profile, mood logs", "Placeholder"],
            ["Document extraction", "Internal tool", "Parse uploads", "Placeholder"],
            ["Direct answer", "No tool", "LLM only", "Built"],
            ["Crisis response", "No tool", "Fixed safety message", "Built"],
        ],
    )

    doc.add_heading("Layer 7: Infrastructure & Observability", level=1)
    add_table(
        doc,
        ["Area", "Status", "What's needed"],
        [
            ["Debug trace panel", "Built", "Add strategy name, gate reasoning, corpus, timing"],
            ["Automated regression checks", "Not built", "Run evals when code changes"],
            ["Unit test coverage", "Not built", "Tests for routing, gating, ranking"],
        ],
    )

    doc.add_heading("Recommended work order", level=1)
    for i, step in enumerate(
        [
            "Consolidate orchestration and gating — single routing system",
            "Lock down gating accuracy — automated regression checks",
            "Pick a clinical retrieval strategy — run bake-off, set default",
            "Implement journal entry retrieval — embed entries, extend search to both corpora",
            "Build response evaluation — AI-as-judge on core test cases",
            "Activate external tools — web search, structured query, document extraction",
            "Experiment with AI-driven orchestration — after rule-based baselines measured",
        ],
        start=1,
    ):
        doc.add_paragraph(f"{i}. {step}")

    doc.add_heading("Quick reference: what exists vs what's left", level=1)
    add_table(
        doc,
        ["Area", "Built", "Tested", "Remaining work"],
        [
            ["Intent routing", "Yes", "Manual only", "Labeled evaluation, unify with gating"],
            ["Crisis handling", "Yes", "Partial", "Dedicated evaluation + automated checks"],
            ["Direct answer vs database search", "Yes", "~200 cases", "Routing strategy A/B, automated eval"],
            ["Corpus selection", "Partial", "Not testable", "Journal retrieval + labels"],
            ["Retrieval — clinical KB", "Yes", "Bake-off", "Pick winner, finer metrics"],
            ["Retrieval — journal", "No", "—", "Build and evaluate"],
            ["Generation / prompting", "Yes", "Manual only", "Prompt testing, model A/B"],
            ["Response quality eval", "No", "—", "AI-as-judge, safety checks"],
            ["External tools", "Placeholder", "—", "Build and evaluate"],
            ["Multi-step agent loop", "No", "—", "Design after single-pass measured"],
            ["Automated eval on code changes", "No", "—", "Wire into dev pipeline"],
        ],
    )

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
